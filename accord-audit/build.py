import collections
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from tracker_data import R, COLS, STATUSES, SECTIONS

ids = [r["ID"] for r in R]
assert len(ids) == len(set(ids)), "duplicate IDs"

wb = Workbook()
ws = wb.active; ws.title = "Tracker"
ws.append(COLS)
for r in R: ws.append([r[c] for c in COLS])
widths = [10,26,44,18,38,80,11,26,30,22,24,52,18,18,44,40]
for i,w in enumerate(widths,1): ws.column_dimensions[chr(64+i)].width = w
hdr = PatternFill("solid", fgColor="0B203B")
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr; c.alignment = Alignment(wrap_text=True, vertical="center")
colors = {"Critical":"F4B6B6","High":"F9D6A5","Medium":"FFF2B3","Low":"D5E8D4","Enhancement":"DDE6F5"}
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    row[6].fill = PatternFill("solid", fgColor=colors[row[6].value])
ws.freeze_panes = "D2"; ws.auto_filter.ref = ws.dimensions
dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
# list validation string must be <255 chars; use a sheet range instead
s2 = wb.create_sheet("Lists")
for i,s in enumerate(STATUSES,1): s2.cell(row=i, column=1, value=s)
dv = DataValidation(type="list", formula1=f"=Lists!$A$1:$A${len(STATUSES)}", allow_blank=False)
ws.add_data_validation(dv); dv.add(f"H2:H{len(R)+1}")

s3 = wb.create_sheet("Summary", 0)
s3.append(["Accord Insurance Brokers master remediation tracker"]); s3["A1"].font = Font(bold=True, size=14)
s3.append(["Read-only audit baseline, 2 Oct 2026. Creating this tracker authorises no website change."]); s3.append([])
s3.append(["Priority","Count"])
pc = collections.Counter(r["Priority"] for r in R)
for p in ["Critical","High","Medium","Low","Enhancement"]: s3.append([p, pc.get(p,0)])
s3.append([]); s3.append(["Status","Count"])
sc = collections.Counter(r["Status"] for r in R)
for s in STATUSES: s3.append([s, sc.get(s,0)])
s3.append([]); s3.append(["Section","Items"])
for n,name in SECTIONS.items(): s3.append([f"{n}. {name}", sum(1 for r in R if r["Section"]==name)])
s3.column_dimensions["A"].width = 58
wb.save("Accord_Master_Remediation_Tracker.xlsx")

# markdown version
md = ["# Accord Insurance Brokers: Master Remediation Tracker", "",
      "Read-only audit baseline, 2 Oct 2026. Creating this tracker authorises no website change. Statuses use only the approved list.", ""]
md += ["| Priority | Count |","|---|---|"] + [f"| {p} | {pc.get(p,0)} |" for p in ["Critical","High","Medium","Low","Enhancement"]] + [""]
for n,name in SECTIONS.items():
    md += [f"## {n}. {name}", ""]
    for r in [x for x in R if x["Section"]==name]:
        md += [f"### {r['ID']}: {r['Finding / task']}", ""]
        for c in COLS:
            if c in ("ID","Section","Finding / task"): continue
            v = r[c]
            if v: md.append(f"- **{c}:** {v}")
        md.append("")
open("Accord_Master_Remediation_Tracker.md","w").write("\n".join(md))
print(len(R), dict(pc)); print(dict(sc))
